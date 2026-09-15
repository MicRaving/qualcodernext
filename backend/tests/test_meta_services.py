"""Unit tests for the meta services (parsers, prompts, normalization)."""

from __future__ import annotations

from pathlib import Path

from qualcoder_api.core.models import MetaCriterion
from qualcoder_api.services.meta_extract import (
    metric_has_full_descriptives,
    metric_has_metafor_fallback,
    normalize_and_validate,
    normalize_metric,
)
from qualcoder_api.services.meta_import import parse_ebsco_xml, parse_hits_xlsx
from qualcoder_api.services.meta_prompt import (
    build_cluster_messages,
    criteria_prompt_hash,
    extract_json,
)
from qualcoder_api.services.meta_screen import assign_entries, parse_verdicts


def _criterion(i: int) -> MetaCriterion:
    return MetaCriterion(
        position=i,
        key=f"criterion{i}",
        label=f"Criterion {i}",
        prompt_text=f"Criterion {i} description.",
    )


def test_extract_json_tolerates_prose_and_fences():
    assert extract_json('```json\n[{"a": 1}]\n```') == [{"a": 1}]
    assert extract_json('prefix {"a": 1} suffix') == {"a": 1}
    assert extract_json("no json here") is None


def test_criteria_prompt_hash_is_deterministic():
    criteria = [_criterion(1), _criterion(2)]
    assert criteria_prompt_hash(criteria) == criteria_prompt_hash(criteria)
    assert criteria_prompt_hash(criteria) != criteria_prompt_hash(criteria[:1])


def test_cluster_messages_include_all_criteria():
    criteria = [_criterion(1), _criterion(2)]
    messages = build_cluster_messages(
        [(0, {"title": "T", "abstract": "A", "doi": "10.x/y"})], criteria
    )
    system = messages[0]["content"]
    assert "Criterion 1 description." in system
    assert "criterion1_applies" in system
    assert "criterion2_certainty" in system
    user = messages[1]["content"]
    assert "T" in user and "10.x/y" in user


def test_parse_ebsco_xml(tmp_path: Path):
    xml = """<?xml version="1.0"?><records>
    <record><title>One</title><abstract>abs</abstract>
    <contributors>Doe, J ; Roe, S</contributors><doi>10.1/a</doi>
    <publicationDate>2019-07-01</publicationDate><source>J</source></record>
    </records>"""
    path = tmp_path / "r.xml"
    path.write_text(xml, encoding="utf-8")
    hits = parse_ebsco_xml(str(path))
    assert len(hits) == 1
    hit = hits[0]
    assert hit["title"] == "One"
    assert hit["authors"] == "Doe, J, Roe, S"
    assert hit["year"] == "2019"
    assert hit["doi"] == "10.1/a"


def test_parse_hits_xlsx(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["title", "abstract", "authors", "doi", "publication year"])
    ws.append(["Two", "abstract two", "A, B", "10.2/b", "2022"])
    ws.append(["No DOI", "x", "C", "", "2023"])
    path = tmp_path / "hits.xlsx"
    wb.save(path)
    hits = parse_hits_xlsx(str(path))
    assert len(hits) == 2
    assert hits[0]["doi"] == "10.2/b"
    assert hits[1]["doi"] is None


def test_parse_hits_xlsx_header_not_on_first_row(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Search results export"])  # title line above the header
    ws.append([])
    ws.append(["Authors", "Title", "Year", "Abstract", "DOI"])
    ws.append(["Doe, J", "Study A", "2020", "abs a", "10.3/a"])
    ws.append(["Roe, S", "Study B", "2021", "abs b", "10.3/b"])
    path = tmp_path / "hits_offset.xlsx"
    wb.save(path)
    hits = parse_hits_xlsx(str(path))
    assert [h["title"] for h in hits] == ["Study A", "Study B"]
    assert hits[0]["authors"] == "Doe, J"
    assert hits[0]["year"] == "2020"


def test_parse_hits_xlsx_german_headers(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Titel", "Zusammenfassung", "Autoren", "DOI", "Jahr"])
    ws.append(["Studie A", "Zusammenfassung A", "Doe, J", "10.4/a", "2020"])
    path = tmp_path / "hits_de.xlsx"
    wb.save(path)
    hits = parse_hits_xlsx(str(path))
    assert len(hits) == 1
    assert hits[0]["title"] == "Studie A"
    assert hits[0]["authors"] == "Doe, J"


def test_parse_hits_xlsx_publication_year_is_year_not_source(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Title", "Authors", "DOI", "Publication Year", "Source"])
    ws.append(["T", "A", "10.5/a", "2019", "Journal X"])
    path = tmp_path / "hits_pub.xlsx"
    wb.save(path)
    hits = parse_hits_xlsx(str(path))
    assert hits[0]["year"] == "2019"
    assert hits[0]["source"] == "Journal X"


def test_normalize_metric_infers_vi_from_d():
    metric = normalize_metric({"d_val": 0.5, "n1i": 20, "n2i": 20})
    assert metric["yi"] == 0.5
    assert metric["vi"] is not None
    assert metric["sei"] is not None
    assert metric_has_metafor_fallback(metric)


def test_normalize_metric_splits_balanced_groups():
    metric = normalize_metric({"n_total": 40, "m1i": 5, "m2i": 6, "sd1i": 1, "sd2i": 1})
    assert metric["n1i"] == 20 and metric["n2i"] == 20
    assert metric_has_full_descriptives(metric)


def test_normalize_and_validate_reports_missing():
    records, missing = normalize_and_validate(
        [{"Study_ID": "X2020_Exp1", "DV_Metrics": [{"dv_name": "dv"}]}]
    )
    assert len(records) == 1
    assert len(missing) == 1
    assert missing[0]["dv_name"] == "dv"


def test_assign_entries_uses_array_order_when_lengths_match():
    parsed = [{"a": 1}, {"a": 2}, {"a": 3}]
    assigned = assign_entries(parsed, 3)
    assert assigned == {0: {"a": 1}, 1: {"a": 2}, 2: {"a": 3}}


def test_assign_entries_maps_one_based_index_when_lengths_differ():
    # Model dropped one study but numbered them 1..N: map by index, not order.
    parsed = [{"index": 1, "a": "first"}, {"index": 3, "a": "third"}]
    assigned = assign_entries(parsed, 3)
    assert assigned == {0: {"index": 1, "a": "first"}, 2: {"index": 3, "a": "third"}}


def test_assign_entries_handles_zero_index_and_out_of_range():
    # Lengths differ, so the index path runs. Index 0 is accepted (0-based) and
    # an out-of-range index falls back to the array position.
    parsed = [{"index": 0, "a": "first"}, {"index": 99, "a": "second"}]
    assigned = assign_entries(parsed, 3)
    assert assigned[0]["a"] == "first"
    assert assigned[1]["a"] == "second"


def test_assign_entries_single_object_cluster():
    assert assign_entries({"index": 1, "a": 1}, 1) == {0: {"index": 1, "a": 1}}


def test_parse_verdicts_normalizes_free_text():
    criteria = [_criterion(1), _criterion(2)]
    entry = {
        "criterion1_applies": "YES",
        "criterion1_certainty": "moderate",
        "criterion2_applies": "maybe",
        "criterion2_certainty": "LOW",
    }
    verdicts = parse_verdicts(entry, criteria)
    assert verdicts["criterion1"] == {"applies": "Yes", "certainty": "Medium"}
    assert verdicts["criterion2"] == {"applies": "Unsure", "certainty": "Low"}


def test_parse_verdicts_accepts_boolean_alt_key():
    criteria = [_criterion(1)]
    verdicts = parse_verdicts({criteria[0].key: True}, criteria)
    assert verdicts[criteria[0].key]["applies"] == "Yes"
    verdicts = parse_verdicts({criteria[0].key: {"applies": "No"}}, criteria)
    assert verdicts[criteria[0].key]["applies"] == "No"

def test_parse_hits_xlsx_skips_cover_sheet(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    cover = wb.active
    cover.title = "Search"
    cover.append(["EBSCOhost search export"])
    cover.append(["Query", '"inoculation"'])
    results = wb.create_sheet("SearchResults")
    results.append(["Title", "Authors", "Abstract", "DOI"])
    results.append(["Study A", "Doe, J", "abs", "10.6/a"])
    path = tmp_path / "hits_sheets.xlsx"
    wb.save(path)
    hits = parse_hits_xlsx(str(path))
    assert [h["title"] for h in hits] == ["Study A"]
