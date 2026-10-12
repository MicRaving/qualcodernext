"""CodingStudio folder importer — hits, screening, scheme, extractions, PDFs.

Reads a reference CodingStudio project folder and mirrors it into the local
``meta_*`` tables so the study can be reviewed, screened, extracted and
analysed inside QCnext.

Folder contract (all paths relative to the folder passed in)::

    Papers/<Key>.pdf                 Text/<Key>.txt
    Extraction/<Key>.json            (JSON list, one object per experiment)
    Codings/<Key> - <coder>.json     (optional extra coder's extractions)
    Screening/<coder>.json           (abstract verdicts keyed by accession no.,
                                      or ``row-<worksheet row>`` when blank)
    Screening/fulltext - <coder>.json (per-coder full-text verdicts, keyed by study key)
    master_dataset.xlsx              (Sheet1: hits + C1..C4/Fulltext_Verdict columns)
    Scripts/coding_web/schema.json   (v3.1 coding schema)

Everything is idempotent: hits dedupe through ``MetaRepository.add_hits``,
screening rows upsert on ``(hit_id, owner, run_id, stage)`` (batched through a
single ``INSERT … ON CONFLICT`` executemany per import), extractions upsert on
``(hit_id, owner, scheme)`` and an already-linked PDF is never re-imported. An
``an -> hit_id`` sidecar under ``<project>/.meta/codingstudio.json`` keeps
re-imports stable even when a hit cannot be re-matched by DOI/title.

The workbook is the reference's base/AI rater: its ``C1..C4`` and
``Intervention`` columns import as the ``LLM/base`` abstract rater, and its
``Fulltext_*`` columns as the ``codingstudio`` full-text rater.
"""

from __future__ import annotations

import asyncio
import contextlib
import json
import logging
import os
import re
from pathlib import Path
from typing import Any

import openpyxl
from sqlalchemy import text
from sqlalchemy.ext.asyncio import async_sessionmaker

from qualcoder_api.core.timeutil import now as _now
from qualcoder_api.persistence.repo.meta_repo import MetaRepository, hit_dedupe_key
from qualcoder_api.services.meta_extract import derive_dv_metrics, records_from_response
from qualcoder_api.services.meta_import import COLUMN_ALIASES
from qualcoder_api.services.meta_roles import (
    LEGACY_ROLE_ALIASES,
    metric_by_role,
    role_map,
    value_by_role,
)
from qualcoder_api.services.meta_scheme import (
    EFFECT_FAMILIES,
    FIELD_TYPES,
    all_fields,
    normalise_codebook,
)

logger = logging.getLogger(__name__)

WORKBOOK_NAME = "master_dataset.xlsx"
SCHEMA_RELATIVE = Path("Scripts") / "coding_web" / "schema.json"
SIDECAR_RELATIVE = Path(".meta") / "codingstudio.json"

#: Cap on the reported ``errors`` list; the true total is ``errors_truncated``
#: more, so a broken folder cannot return tens of thousands of strings.
MAX_ERRORS = 50

#: The four inclusion criteria plus the informational intervention-type
#: classification, worded as in the reference ``screening_lib.py``.
_CRITERION_SUFFIXES: tuple[tuple[str, str, str], ...] = (
    (
        "C1",
        "Misinformation intervention",
        "Does the study feature an intervention aimed at reducing susceptibility to "
        "misinformation (inoculation/prebunking, debunking, media literacy, "
        "source-credibility or lateral-reading training, accuracy nudges, social norms, "
        "platform exposure reduction)?",
    ),
    (
        "C2",
        "Control group",
        "Does the study include a control, placebo, or comparison group that does NOT "
        "receive the misinformation intervention?",
    ),
    ("C3", "Language (English)", "Is the paper in English?"),
    (
        "C4",
        "Experimental (primary quantitative study)",
        "Is this a primary empirical study with a quantitative experimental or "
        "quasi-experimental design?",
    ),
)

#: Fallback intervention types when the folder carries no ``schema.json``.
#: The retired ``none``/``generic`` values are intentionally absent: the
#: reference dropped them, keeping the 14 current types (13 + ``other``).
DEFAULT_INTERVENTION_TYPES: tuple[str, ...] = (
    "inoculation",
    "prebunking_warning",
    "lateral_reading",
    "credibility_labels",
    "media_literacy_tips",
    "fact_checking_training",
    "openmindedness_intervention",
    "debunking_and_rebuttals",
    "community_correction",
    "accuracy_prompts",
    "friction",
    "social_norms",
    "platform_exposure_reduction",
    "other",
)

#: Field names the reference stores as arrays of objects.
_OBJECTLIST_FIELDS = frozenset(
    {"Covariates", "factor_levels", "cell_descriptives"}
)

#: Free-text / identifier fields that must not flood the moderator catalogue.
#: Matched by exact name, prefix or suffix (case-sensitive names, as coded).
_FREE_TEXT_FIELDS = frozenset({"Title", "Authors", "DOI", "Study_ID"})
_FREE_TEXT_PREFIXES = ("Author_", "Affiliation_", "ORCID", "source_quote")
_FREE_TEXT_SUFFIXES = ("_Note", "_Notes", "_Reason")


def _is_moderator(name: str) -> bool:
    """False for free-text/identifier fields (see the task's moderator hygiene)."""
    if name in _FREE_TEXT_FIELDS:
        return False
    if name.startswith(_FREE_TEXT_PREFIXES):
        return False
    if name.endswith(_FREE_TEXT_SUFFIXES):
        return False
    return "_rationale" not in name.lower()


#: Field name -> (vocabulary key, editor) for fields whose controlled option
#: list lives in ``schema.json`` ``vocabularies`` rather than inline.
_VOCAB_EDITORS: dict[str, tuple[str, str]] = {
    "dv_domain": ("dv_domains", "multiselect"),
    "score_types": ("score_types", "multiselect"),
    "contrast_type": ("contrast_types", "select"),
    "posttest_type": ("posttest_types", "select"),
    "Publication": ("publication_types", "select"),
    "Publication_Type": ("publication_types", "select"),
    "Publication_Category": ("publication_types", "select"),
    "effect_size_type_reported": ("effect_size_types", "select"),
    "Intervention_Type": ("intervention_types", "multiselect"),
    "Technique_Taxonomy": ("technique_taxonomy", "multiselect"),
    "Attention_Cues": ("attention_cues", "multiselect"),
    "Media_Support": ("media_support", "multiselect"),
    "Comm_Strategies": ("comm_strategies", "multiselect"),
    "Practice_Strategy": ("practice_strategy", "multiselect"),
    "Intervention_Components": ("intervention_components", "multiselect"),
}

#: ``LEGACY_ROLE_ALIASES`` inverted: exact reference field name -> role.
_EXACT_ROLE: dict[str, str] = {
    name: role for role, names in LEGACY_ROLE_ALIASES.items() for name in names
}

#: Effect family -> the roles whose presence enables it.
_FAMILY_ROLES: dict[str, set[str]] = {
    "descriptives": {
        "n1", "m1", "sd1", "n2", "m2", "sd2",
        "n1_pre", "m1_pre", "sd1_pre", "n2_pre", "m2_pre", "sd2_pre",
        "n1_post", "m1_post", "sd1_post", "n2_post", "m2_post", "sd2_post",
        "n_total",
    },
    "t": {"t"},
    "F": {"F"},
    "r": {"r"},
    "yi": {"yi", "vi", "sei"},
    "d_g": {"d", "g", "d_ci_low", "d_ci_high", "g_ci_low", "g_ci_high", "d_se", "g_se"},
    "odds_ratio": {"or", "or_ci_low", "or_ci_high", "log_odds"},
    "risk_ratio": {"rr", "rr_ci_low", "rr_ci_high"},
    "beta": {"beta", "beta_se"},
    "eta2p": {"eta2p", "eta2p_ci_low", "eta2p_ci_high"},
    "chi2": {"chi2"},
}

_FULLTEXT_KEY_RE = re.compile(r"\bC([1-4])\b", re.IGNORECASE)


# ── small text helpers ────────────────────────────────────────────────────

def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _clean_optional(value: Any) -> str | None:
    text = _clean(value)
    return text or None


def _normalise_doi(value: Any) -> str:
    text = _clean(value)
    lowered = text.lower()
    for prefix in ("https://doi.org/", "http://doi.org/", "https://dx.doi.org/", "doi:"):
        if lowered.startswith(prefix):
            text = text[len(prefix) :]
            break
    return text.strip()


def _normalise_key(value: Any) -> str:
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def _identity_key(doi: Any, title: Any, authors: Any) -> str | None:
    return hit_dedupe_key(doi=_clean_optional(doi), title=_clean_optional(title), authors=_clean_optional(authors))


def _role_or_raw(record: Any, roles: dict[str, str] | None, role: str) -> Any:
    if not isinstance(record, dict):
        return None
    if roles:
        found = value_by_role(record, roles, role)
        if found not in (None, ""):
            return found
    return metric_by_role(record).get(role)


# ── folder scan ───────────────────────────────────────────────────────────

def _safe_glob(root: Path, pattern: str) -> list[Path]:
    with contextlib.suppress(OSError):
        return sorted(p for p in root.glob(pattern) if p.is_file())
    return []


def _scan_folder(folder: str) -> dict[str, Any]:
    root = Path(folder)
    return {
        "root": root,
        "papers": _safe_glob(root, "Papers/*.pdf"),
        "texts": _safe_glob(root, "Text/*.txt"),
        "extractions": _safe_glob(root, "Extraction/*.json"),
        "codings": _safe_glob(root, "Codings/*.json"),
        "screenings": _safe_glob(root, "Screening/*.json"),
        "workbook": root / WORKBOOK_NAME,
        "schema": root / SCHEMA_RELATIVE,
    }


def _read_json(path: Path) -> Any:
    with open(path, encoding="utf-8") as handle:
        return json.load(handle)


def _record_count(path: Path) -> int:
    """Number of experiment objects in one ``Extraction/<Key>.json`` list."""
    data = _read_json(path)
    if isinstance(data, list):
        return len([r for r in data if isinstance(r, dict)])
    if isinstance(data, dict):
        return 1
    return 0


def _screening_files(paths: list[Path]) -> list[tuple[str, str, Path]]:
    """``(coder, stage, path)`` for each ``Screening/*.json`` file."""
    out: list[tuple[str, str, Path]] = []
    for path in paths:
        stem = path.stem
        lowered = stem.lower()
        if lowered in ("fulltext", "fulltext_agreed", "fulltext-agreed"):
            continue
        if lowered.startswith(("fulltext - ", "fulltext-")):
            coder = stem.split("-", 1)[1].strip()
            out.append((coder or "coder", "fulltext", path))
        else:
            out.append((stem.strip() or "coder", "abstract", path))
    return out


# ── criteria ──────────────────────────────────────────────────────────────

def _criteria_payload(intervention_options: list[str] | None = None) -> list[dict[str, Any]]:
    """The four gating criteria plus the informational intervention classify."""
    criteria: list[dict[str, Any]] = []
    for position, (key, label, prompt) in enumerate(_CRITERION_SUFFIXES):
        criteria.append(
            {
                "position": position,
                "key": key,
                "label": label,
                "prompt_text": prompt,
                "kind": "verdict",
                "options": ["Yes", "No", "Unsure"],
                "pass_values": ["Yes"],
                "exclude_reason": f"{label} ({key})",
                "gating": 1,
                "help": prompt,
            }
        )
    options: list[str] = list(intervention_options or DEFAULT_INTERVENTION_TYPES)
    criteria.append(
        {
            "position": len(criteria),
            "key": "Intervention",
            "label": "Intervention type",
            "prompt_text": (
                "Which misinformation intervention family does the study evaluate? "
                "Informational only — never blocks inclusion."
            ),
            "kind": "classify",
            "options": options,
            "pass_values": ["Yes"],
            "exclude_reason": "",
            "gating": 0,
            "help": "",
        }
    )
    return criteria


def _intervention_options(schema: dict[str, Any] | None) -> list[str]:
    if isinstance(schema, dict):
        vocab = schema.get("vocabularies")
        if isinstance(vocab, dict):
            types = vocab.get("intervention_types")
            if isinstance(types, list):
                seen: set[str] = set()
                out: list[str] = []
                for value in types:
                    text = _clean(value)
                    if text and text not in seen:
                        seen.add(text)
                        out.append(text)
                if out:
                    return out
    return list(DEFAULT_INTERVENTION_TYPES)


# ── workbook parsing ──────────────────────────────────────────────────────

def _detect_columns(header: list[str]) -> dict[str, int]:
    """Map canonical hit fields to columns using the shared import aliases."""
    cells = [str(cell or "").lower().strip() for cell in header]
    mapping: dict[str, int] = {}
    used: set[int] = set()
    for exact in (True, False):
        for canonical, aliases in COLUMN_ALIASES.items():
            if canonical in mapping:
                continue
            for idx, cell in enumerate(cells):
                if idx in used or not cell:
                    continue
                matched = cell == canonical or cell in aliases
                if not matched and not exact:
                    matched = any(alias in cell for alias in aliases)
                if matched:
                    mapping[canonical] = idx
                    used.add(idx)
                    break
    return mapping


def _read_workbook(path: Path) -> tuple[list[str], list[tuple[Any, ...]]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = ["" if cell is None else str(cell) for cell in next(rows)]
        return header, list(rows)
    finally:
        workbook.close()


def _count_workbook_rows(path: Path) -> int:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.worksheets[0]
        return max(0, int(sheet.max_row or 0) - 1)
    finally:
        workbook.close()


def _workbook_studies(path: Path) -> list[dict[str, Any]]:
    """One entry per usable workbook row (hit fields + screening verdicts).

    The workbook is the reference's base/AI rater: its ``C1..C4`` and
    ``Intervention`` verdict/certainty/rationale columns are captured as the
    ``abstract`` block and its ``Fulltext_Verdict``/``Fulltext_*`` columns as
    the ``fulltext`` block. ``decision`` holds the row's explicit
    ``Effective_Verdict`` (falling back to ``Fulltext_Verdict``).
    """
    header, rows = _read_workbook(path)
    columns = _detect_columns(header)
    by_name = {str(name).strip().lower(): idx for idx, name in enumerate(header)}

    def cell(row: tuple[Any, ...], field: str) -> str:
        idx = columns.get(field)
        if idx is None or idx >= len(row):
            return ""
        return _clean(row[idx])

    def named(row: tuple[Any, ...], name: str) -> str:
        idx = by_name.get(name.lower())
        if idx is None or idx >= len(row):
            return ""
        return _clean(row[idx])

    def criterion_block(row: tuple[Any, ...]) -> dict[str, Any] | None:
        criteria: dict[str, Any] = {}
        for lab in ("C1", "C2", "C3", "C4"):
            verdict = named(row, f"{lab}_verdict")
            if verdict:
                criteria[lab] = {
                    "verdict": verdict,
                    "certainty": named(row, f"{lab}_certainty"),
                    "rationale": named(row, f"{lab}_rationale"),
                }
        intervention = named(row, "Intervention_verdict")
        if intervention:
            criteria["Intervention"] = {
                "verdict": intervention,
                "certainty": named(row, "Intervention_certainty"),
                "rationale": named(row, "Intervention_rationale"),
            }
        return {"criteria": criteria} if criteria else None

    studies: list[dict[str, Any]] = []
    for index, row in enumerate(rows):
        title = cell(row, "title")
        doi = _normalise_doi(cell(row, "doi"))
        abstract = cell(row, "abstract")
        if not (title or doi or abstract):
            continue
        accession = named(row, "an")
        # The reference keys a workbook row without an accession as
        # ``row-<worksheet row>`` (1-based: the header is row 1, so the first
        # data row is 2) — see ``screening_lib.py``'s ``uid = an or f"row-{n+2}"``.
        key = accession or f"row-{index + 2}"
        fulltext = None
        verdict = named(row, "Fulltext_Verdict")
        if verdict:
            fulltext = {
                "verdict": verdict,
                "criterion": named(row, "Fulltext_Exclusion_Criterion"),
                "reason": named(row, "Fulltext_Reason"),
            }
        studies.append(
            {
                "key": key,
                "an": accession,
                "title": title,
                "doi": doi or None,
                "authors": cell(row, "authors") or None,
                "decision": named(row, "Effective_Verdict") or verdict,
                "abstract": criterion_block(row),
                "hit": {
                    "title": title,
                    "abstract": abstract or None,
                    "authors": cell(row, "authors") or None,
                    "doi": doi or None,
                    "year": cell(row, "year") or None,
                    "source": cell(row, "source") or None,
                    "language": cell(row, "language") or None,
                },
                "fulltext": fulltext,
            }
        )
    return studies


# ── coding schema -> codebook v2 ──────────────────────────────────────────

def _schema_type(raw: Any) -> str:
    text = _clean(raw).split("|")[0].strip().lower()
    if text == "array":
        text = "list"
    return text if text in FIELD_TYPES else "string"


def _default_editor(type_name: str) -> str:
    return {
        "integer": "number",
        "number": "number",
        "boolean": "boolean",
        "list": "multiselect",
        "object": "json",
    }.get(type_name, "text")


def _schema_editor(raw: dict[str, Any], type_name: str) -> tuple[str, list[str]]:
    vocab = raw.get("vocab")
    if isinstance(vocab, list) and vocab:
        return "multiselect", [_clean(v) for v in vocab if _clean(v)]
    options = raw.get("options")
    if isinstance(options, list) and options:
        labels: list[str] = []
        for option in options:
            if isinstance(option, dict):
                label = _clean(option.get("label")) or _clean(option.get("value"))
            else:
                label = _clean(option)
            if label:
                labels.append(label)
        if labels:
            return "select", labels
    editor = _clean(raw.get("editor")).lower()
    if editor == "chips":
        return "multiselect", []
    if editor:
        return editor, []
    return _default_editor(type_name), []


def _ordered_vocab_items(vocab: dict[Any, Any]) -> list[tuple[Any, Any]]:
    """Vocabulary items sorted by numeric key when the keys are numeric codes."""
    def sort_key(item: tuple[Any, Any]) -> tuple[int, Any]:
        key = str(item[0])
        return (0, int(key)) if key.isdigit() else (1, key)

    return sorted(vocab.items(), key=sort_key)


def _inline_label(raw: dict[str, Any], key: Any) -> str | None:
    """The schema field's label for a coded value, when it declares one."""
    options = raw.get("options")
    if not isinstance(options, list):
        return None
    for option in options:
        if isinstance(option, dict):
            if str(option.get("value")) == str(key):
                return _clean(option.get("label")) or _clean(option.get("value"))
        elif str(option) == str(key):
            return _clean(option)
    return None


def _vocab_option_list(vocab: Any, raw: dict[str, Any]) -> list[str]:
    """Option labels for a vocabulary entry.

    List vocabularies contribute their values directly; numeric-keyed dict
    vocabularies (``contrast_types`` …) contribute one option per code, using
    the schema field's own label for that code when it declares one so the
    editor stays human-readable.
    """
    out: list[str] = []
    if isinstance(vocab, list):
        for value in vocab:
            text = _clean(value)
            if text:
                out.append(_inline_label(raw, text) or text)
    elif isinstance(vocab, dict):
        for key, value in _ordered_vocab_items(vocab):
            out.append(_inline_label(raw, key) or _clean(value) or str(key))
    seen: set[str] = set()
    unique: list[str] = []
    for option in out:
        if option and option not in seen:
            seen.add(option)
            unique.append(option)
    return unique


def _level_option_list(levels: Any) -> list[str]:
    """Option labels from a ``module_field_levels`` entry (values are labels)."""
    if isinstance(levels, dict):
        return [_clean(value) for _, value in _ordered_vocab_items(levels) if _clean(value)]
    if isinstance(levels, list):
        return [_clean(value) for value in levels if _clean(value)]
    return []


def _schema_field(raw: Any, default_group: str) -> dict[str, Any] | None:
    if not isinstance(raw, dict):
        return None
    name = _clean(raw.get("name"))
    if not name:
        return None
    type_name = _schema_type(raw.get("type"))
    editor: str
    options: list[str]
    if name in _OBJECTLIST_FIELDS:
        type_name = "list"
        editor, options = "objectlist", []
    else:
        editor, options = _schema_editor(raw, type_name)
    module = _clean(raw.get("module")) or None
    return {
        "name": name,
        "label": _clean(raw.get("label")) or name,
        "group": _clean(raw.get("group")) or default_group,
        "type": type_name,
        "editor": editor,
        "options": options,
        "columns": raw.get("columns") if isinstance(raw.get("columns"), list) else [],
        "required": bool(raw.get("required")),
        "description": _clean(raw.get("description")),
        "role": _EXACT_ROLE.get(name),
        "module": module,
        #: Per-module option lists for the 7 field names that repeat across
        #: modules with different levels (see ``schema_codebook``).
        "module_options": {},
        "moderator": _is_moderator(name),
    }


def _effect_families(metric_fields: list[dict[str, Any]]) -> list[str]:
    roles = {field.get("role") for field in metric_fields if field.get("role")}
    return [family for family in EFFECT_FAMILIES if roles & _FAMILY_ROLES.get(family, set())]


def schema_codebook(schema: dict[str, Any]) -> dict[str, Any]:
    """Convert a CodingStudio v3.1 ``schema.json`` to a codebook v2 dict.

    **Module identity and collisions.** Each module field keeps its ``module``
    and the codebook records ``intervention_type_modules`` so a reader can map
    a record's ``Intervention_Type`` to the modules that apply. Seven field
    names repeat across modules with different option lists
    (``Assessment_Paradigm``, ``Intervention_Source_Concept``, ``Human_Trainer``,
    ``Correction_Source``, ``Timing_Relative_to_Misinformation``,
    ``Primary_Outcome``, ``Intervention_Delivery``). The field is emitted
    **once** with the first occurrence's options (renaming it per module would
    break the importer, which merges study data by field name), and every
    module's own list is attached under ``module_options[<module>]`` — sourced
    from ``vocabularies.module_field_levels`` and from the duplicate field
    declarations — so an editor can present the options for the record's
    module while the importer stays correct.

    **Vocabularies.** Controlled option lists declared in ``vocabularies`` are
    overlaid onto their fields (see ``_VOCAB_EDITORS``); the retired
    ``none``/``generic`` intervention types are omitted.
    """
    groups = [_clean(g) for g in (schema.get("groups") or []) if _clean(g)]
    vocabularies = schema.get("vocabularies")
    vocabularies = vocabularies if isinstance(vocabularies, dict) else {}
    module_levels = vocabularies.get("module_field_levels")
    module_levels = module_levels if isinstance(module_levels, dict) else {}

    raw_by_name: dict[str, dict[str, Any]] = {}
    study_by_name: dict[str, dict[str, Any]] = {}
    metric_by_name: dict[str, dict[str, Any]] = {}

    def collect(raws: Any, default_group: str, out: list[dict[str, Any]],
                seen: dict[str, dict[str, Any]]) -> None:
        for raw in raws or []:
            field = _schema_field(raw, default_group)
            if field is None:
                continue
            name = field["name"]
            if isinstance(raw, dict):
                raw_by_name.setdefault(name, raw)
            module = field.get("module") or default_group
            duplicate = seen.get(name)
            if duplicate is not None:
                if field["options"]:
                    duplicate.setdefault("module_options", {})[module] = field["options"]
                continue
            seen[name] = field
            out.append(field)

    study_fields_out: list[dict[str, Any]] = []
    collect(schema.get("common_fields") or [], "Study", study_fields_out, study_by_name)
    modules = schema.get("modules")
    if isinstance(modules, dict):
        for module in modules.values():
            fields = (module or {}).get("fields") if isinstance(module, dict) else None
            collect(fields or [], "Intervention module", study_fields_out, study_by_name)

    metric_fields_out: list[dict[str, Any]] = []
    collect(schema.get("dv_fields") or [], "DV_Metrics", metric_fields_out, metric_by_name)

    # Attach the authoritative per-module levels to the fields that repeat.
    for name, levels in module_levels.items():
        field = study_by_name.get(name) or metric_by_name.get(name)
        if field is None or not isinstance(levels, dict):
            continue
        overrides = field.setdefault("module_options", {})
        for module, codes in levels.items():
            options = _level_option_list(codes)
            if options:
                overrides[str(module)] = options

    # Overlay the controlled vocabularies.
    for field in [*study_fields_out, *metric_fields_out]:
        info = _VOCAB_EDITORS.get(field["name"])
        if info is None:
            continue
        vocab_key, editor = info
        vocab = vocabularies.get(vocab_key)
        if vocab is None:
            continue
        options = _vocab_option_list(vocab, raw_by_name.get(field["name"], {}))
        if options:
            field["editor"] = editor
            field["options"] = options

    itm = schema.get("intervention_type_modules")
    return {
        "version": 2,
        "category": "Meta",
        "groups": groups or ["Study", "DV_Metrics"],
        "study_fields": study_fields_out,
        "fields": study_fields_out,
        "metrics": {
            "group": "DV_Metrics",
            "label": "Effect sizes",
            "fields": metric_fields_out,
        },
        "effect_families": _effect_families(metric_fields_out),
        "intervention_type_modules": (
            {str(k): [str(m) for m in v] for k, v in itm.items() if isinstance(v, list)}
            if isinstance(itm, dict)
            else {}
        ),
    }


def minimal_codebook() -> dict[str, Any]:
    """Small v1 codebook used when the folder carries no ``schema.json``."""
    study = [
        {"name": "Study_ID", "type": "string", "editor": "text", "role": "study_id"},
        {"name": "Title", "type": "string", "editor": "text", "role": "title"},
        {"name": "Authors", "type": "string", "editor": "text", "role": "authors"},
        {"name": "DOI", "type": "string", "editor": "text", "role": "doi"},
        {"name": "Year", "type": "integer", "editor": "number", "role": "year"},
        {"name": "N_Analyzed", "type": "integer", "editor": "number", "role": "sample_n"},
    ]
    metric = [
        {"name": "dv_name", "type": "string", "editor": "text", "role": "dv_name"},
        {"name": "n1i", "type": "number", "editor": "number", "role": "n1"},
        {"name": "n2i", "type": "number", "editor": "number", "role": "n2"},
        {"name": "m1i", "type": "number", "editor": "number", "role": "m1"},
        {"name": "m2i", "type": "number", "editor": "number", "role": "m2"},
        {"name": "sd1i", "type": "number", "editor": "number", "role": "sd1"},
        {"name": "sd2i", "type": "number", "editor": "number", "role": "sd2"},
        {"name": "yi", "type": "number", "editor": "number", "role": "yi"},
        {"name": "vi", "type": "number", "editor": "number", "role": "vi"},
        {"name": "d_val", "type": "number", "editor": "number", "role": "d"},
    ]
    return {
        "version": 1,
        "groups": ["Study", "DV_Metrics"],
        "study_fields": study,
        "metrics": {"group": "DV_Metrics", "label": "Effect sizes", "fields": metric},
    }


# ── preview ───────────────────────────────────────────────────────────────

def _sample_studies(paths: list[Path], roles: dict[str, str] | None = None) -> list[dict[str, Any]]:
    samples: list[dict[str, Any]] = []
    for path in paths[:5]:
        try:
            data = _read_json(path)
        except (OSError, ValueError):
            continue
        records = data if isinstance(data, list) else [data]
        records = [r for r in records if isinstance(r, dict)]
        first = records[0] if records else {}
        samples.append(
            {
                "key": path.stem,
                "title": _clean(_role_or_raw(first, roles, "title")),
                "doi": _clean(_role_or_raw(first, roles, "doi")),
                "records": len(records),
            }
        )
    return samples


def preview_codingstudio(folder: str) -> dict[str, Any]:
    """Read-only description of a CodingStudio folder (never mutates it)."""
    scan = _scan_folder(folder)
    root: Path = scan["root"]
    errors: list[str] = []
    exists = root.is_dir()

    study_count = 0
    for path in scan["extractions"]:
        try:
            study_count += _record_count(path)
        except (OSError, ValueError) as err:
            errors.append(f"{path.name}: {err}")

    screenings: list[dict[str, Any]] = []
    for coder, _stage, path in _screening_files(scan["screenings"]):
        try:
            data = _read_json(path)
            rows = len(data) if isinstance(data, dict) else len(data or [])
        except (OSError, ValueError) as err:
            errors.append(f"{path.name}: {err}")
            rows = 0
        screenings.append({"coder": coder, "rows": rows})

    workbook = None
    if scan["workbook"].is_file():
        try:
            workbook = {"rows": _count_workbook_rows(scan["workbook"]), "path": str(scan["workbook"])}
        except (OSError, ValueError) as err:
            errors.append(f"{WORKBOOK_NAME}: {err}")

    schema_version: str | None = None
    if scan["schema"].is_file():
        try:
            schema = _read_json(scan["schema"])
            if isinstance(schema, dict):
                schema_version = _clean_optional(schema.get("schema_version"))
        except (OSError, ValueError) as err:
            errors.append(f"schema.json: {err}")

    return {
        "folder": str(folder),
        "exists": exists,
        "papers": len(scan["papers"]),
        "texts": len(scan["texts"]),
        "extraction_files": len(scan["extractions"]),
        "study_count": study_count,
        "screenings": screenings,
        "workbook": workbook,
        "schema_version": schema_version,
        "sample_studies": _sample_studies(scan["extractions"]) if exists else [],
        "errors": errors,
    }


# ── screening verdict builders ────────────────────────────────────────────

def _match_fulltext_criterion(criterion: str, reason: str, criteria: list[dict[str, Any]]) -> str | None:
    text = f"{criterion} {reason}".casefold()
    match = _FULLTEXT_KEY_RE.search(text)
    if match:
        return f"C{match.group(1)}"
    for item in criteria:
        if not item.get("gating"):
            continue
        label = str(item.get("label") or "").casefold()
        if label and label in text:
            return str(item["key"])
    return None


def _fulltext_verdicts(
    entry: dict[str, Any], criteria: list[dict[str, Any]]
) -> tuple[dict[str, Any], str]:
    """``({criterion: verdict}, row_reason)`` for one full-text decision."""
    exclude = entry.get("exclude")
    if exclude is None:
        exclude = not str(entry.get("verdict") or "").lower().startswith("inc")
    exclude = bool(exclude)
    criterion = _clean(entry.get("criterion") or entry.get("criterion_label"))
    reason = _clean(entry.get("reason") or entry.get("note") or entry.get("criterion_label"))
    failing = _match_fulltext_criterion(criterion, reason, criteria) if exclude else None

    verdicts: dict[str, Any] = {}
    for item in criteria:
        if not item.get("gating"):
            continue
        applies = "Yes"
        if exclude:
            applies = "No" if failing is None or item["key"] == failing else "Yes"
        row: dict[str, Any] = {"applies": applies, "certainty": "Medium"}
        if applies == "No":
            row["reason"] = reason or criterion
        verdicts[str(item["key"])] = row
    return verdicts, reason


def _abstract_verdicts(entry: dict[str, Any]) -> tuple[dict[str, Any], str | None, str]:
    """``(verdicts, intervention_type, rationale)`` from a coder verdict block."""
    criteria = entry.get("criteria") if isinstance(entry, dict) else None
    verdicts: dict[str, Any] = {}
    intervention: str | None = None
    rationales: list[str] = []
    if isinstance(criteria, dict):
        for key, value in criteria.items():
            if not isinstance(value, dict):
                continue
            applies = value.get("verdict")
            if applies is None:
                applies = value.get("applies")
            row: dict[str, Any] = {
                "applies": "" if applies is None else str(applies),
                "certainty": _clean(value.get("certainty")) or "Medium",
            }
            rationale = _clean(value.get("rationale") or value.get("reason"))
            if rationale:
                row["reason"] = rationale
                rationales.append(rationale)
            verdicts[str(key)] = row
            if str(key).strip().lower() == "intervention" and applies not in (None, ""):
                intervention = str(applies)
    rationale = " ".join(dict.fromkeys(rationales))
    return verdicts, intervention, rationale


# ── hits, identity and the an -> hit_id sidecar ───────────────────────────

def _hits_lookup(hits: list[Any]) -> dict[str, int]:
    lookup: dict[str, int] = {}
    for hit in hits:
        key = hit_dedupe_key(doi=hit.doi, title=hit.title, authors=hit.authors)
        if key and key not in lookup:
            lookup[key] = hit.hit_id
    return lookup


def _title_lookup(hits: list[Any]) -> dict[str, int]:
    """Secondary fallback keyed by the normalized title alone.

    Reference workbooks contain rows with neither a DOI nor authors, so
    ``hit_dedupe_key`` yields nothing and the primary lookup cannot find the
    hit the row was inserted as; the title is the only remaining handle.
    """
    lookup: dict[str, int] = {}
    for hit in hits:
        key = _normalise_key(hit.title or "")
        if key and key not in lookup:
            lookup[key] = hit.hit_id
    return lookup


def _records_identity(
    records: list[dict[str, Any]], roles: dict[str, str] | None
) -> tuple[Any, Any, Any, Any]:
    study = records[0].get("study") if records and isinstance(records[0], dict) else {}
    if not isinstance(study, dict):
        study = {}
    return (
        _role_or_raw(study, roles, "doi"),
        _role_or_raw(study, roles, "title"),
        _role_or_raw(study, roles, "authors"),
        _role_or_raw(study, roles, "year"),
    )


def _studies_from_extractions(paths: list[Path], codebook: dict[str, Any]) -> list[dict[str, Any]]:
    """Fallback: one hit per ``Extraction/<Key>.json`` (no workbook)."""
    roles = role_map(all_fields(codebook))
    studies: list[dict[str, Any]] = []
    for path in paths:
        data = _read_json(path)
        records, _ = records_from_response(data, codebook)
        doi, title, authors, _year = _records_identity(records, roles)
        study = records[0].get("study") if records and isinstance(records[0], dict) else {}
        year = _role_or_raw(study if isinstance(study, dict) else {}, roles, "year")
        hit = {
            "title": _clean(title) or path.stem,
            "abstract": None,
            "authors": _clean_optional(authors),
            "doi": _clean_optional(doi),
            "year": _clean_optional(year),
            "source": None,
            "language": None,
        }
        studies.append(
            {
                "key": path.stem,
                "an": "",
                "title": hit["title"],
                "doi": hit["doi"],
                "authors": hit["authors"],
                "decision": "",
                "abstract": None,
                "hit": hit,
                "fulltext": None,
            }
        )
    return studies


def _sidecar_path(project_path: str | Path) -> Path:
    return Path(project_path) / SIDECAR_RELATIVE


def _load_sidecar(project_path: str | Path) -> dict[str, int]:
    path = _sidecar_path(project_path)
    if not path.is_file():
        return {}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    hits = data.get("hits") if isinstance(data, dict) else None
    if not isinstance(hits, dict):
        return {}
    out: dict[str, int] = {}
    for key, value in hits.items():
        with contextlib.suppress(TypeError, ValueError):
            out[str(key)] = int(value)
    return out


def _save_sidecar(project_path: str | Path, mapping: dict[str, int], folder: str) -> None:
    path = _sidecar_path(project_path)
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "folder": str(folder),
            "updated": _now(),
            "hits": {str(k): int(v) for k, v in mapping.items()},
        }
        path.write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    except OSError as err:  # pragma: no cover - best-effort bookkeeping
        logger.warning("Could not write CodingStudio sidecar %s: %s", path, err)


def _resolve_studies(
    studies: list[dict[str, Any]],
    lookup: dict[str, int],
    title_lookup: dict[str, int],
    sidecar: dict[str, int],
    existing_ids: set[int],
) -> dict[str, int]:
    resolved: dict[str, int] = {}
    for study in studies:
        key = str(study["key"])
        hit_id = sidecar.get(key)
        if hit_id not in existing_ids:
            identity = _identity_key(study.get("doi"), study.get("title"), study.get("authors"))
            hit_id = lookup.get(identity) if identity else None
        if hit_id is None:
            hit_id = title_lookup.get(_normalise_key(study.get("title")))
        if hit_id is not None:
            resolved[key] = hit_id
    return resolved


# ── PDF linking ───────────────────────────────────────────────────────────

async def _store_pdf(
    project_path: str,
    session_factory: async_sessionmaker,
    repo: MetaRepository,
    hit: Any,
    *,
    content: bytes,
    filename: str,
) -> tuple[int, int]:
    """Import PDF bytes as a source and link them to ``hit`` (manual method).

    Mirrors the API's ``_store_pdf`` helper so the importer owns no request.
    """
    from qualcoder_api.core.security import sanitize_filename
    from qualcoder_api.services.import_service import ImportService
    from qualcoder_api.services.user_settings import get_codername

    safe_name = sanitize_filename(filename or "paper.pdf", "meta_pdf")
    tmp = os.path.join(project_path, f"_meta_codingstudio_{hit.hit_id}_{safe_name}")
    with open(tmp, "wb") as out:  # noqa: ASYNC230 - small local temp write
        out.write(content)
    try:
        importer = ImportService(project_path, session_factory)
        source = await importer.import_file(
            tmp, owner=get_codername(), filename=f"meta_{hit.hit_id}_{safe_name}"
        )
        if source is None:
            raise ValueError("file import failed (duplicate name?)")
        doc = await repo.upsert_document(
            hit_id=hit.hit_id,
            source_id=source.id,
            retrieval_method="manual",
            url=hit.doi or "",
            status="done",
            message="imported from CodingStudio",
        )
        await repo.set_hit_status(hit.hit_id, "downloaded")
        return source.id, doc.id
    finally:
        with contextlib.suppress(OSError):
            os.remove(tmp)


def _first_author_key(authors: Any) -> str:
    """Normalized first-author key from an author string (``Doe, J; Roe, S``)."""
    if not authors:
        return ""
    return _normalise_key(re.split(r"[;,]", str(authors))[0])


def _author_year_key(authors: Any, year: Any) -> str:
    """``<first author>|<year>`` key, or ``""`` when either half is missing."""
    author = _first_author_key(authors)
    yr = _clean(year)
    if not author or not yr:
        return ""
    return f"{author}|{yr}"


def _stem_author_year_key(stem: str) -> str:
    """``Smith et al. (2023)`` -> ``smith|2023`` (first author + year).

    The reference names ``Papers/*.pdf`` ``Author (Year).pdf`` (with ``et al.``
    or ``&`` for co-authors), so the file name is a usable last-resort key when
    the extraction record carries neither a DOI nor a matching title.
    """
    match = re.match(r"^(.*?)\s*\((\d{4})[^)]*\)\s*$", stem)
    if not match:
        return ""
    name = re.sub(r"\s+et al\.?$", "", match.group(1).strip(), flags=re.IGNORECASE)
    name = re.split(r"\s*(?:&|,|\band\b)\s*", name)[0].strip()
    return _author_year_key(name, match.group(2))


def _pdf_matcher(
    hits: list[Any],
    resolved: dict[str, int],
    title_lookup: dict[str, int],
    extraction_identity: dict[str, tuple[Any, Any, Any, Any]],
):
    """``match(stem) -> (hit | None, reason)`` for a ``Papers/*.pdf``.

    Matching runs in fidelity order: the exact study key, the extraction
    record's DOI, its normalized title, its first-author+year, and finally the
    ``Author (Year)`` file name against the hits' first-author+year.
    """
    by_id = {hit.hit_id: hit for hit in hits}
    by_doi: dict[str, int] = {}
    by_author_year: dict[str, int] = {}
    for hit in hits:
        doi = _normalise_key(_normalise_doi(hit.doi)) if hit.doi else ""
        if doi and doi not in by_doi:
            by_doi[doi] = hit.hit_id
        key = _author_year_key(hit.authors, hit.year)
        if key and key not in by_author_year:
            by_author_year[key] = hit.hit_id

    def match(stem: str) -> tuple[Any | None, str]:
        hit_id = resolved.get(stem)
        if hit_id in by_id:
            return by_id[hit_id], "study key"
        doi, title, authors, year = extraction_identity.get(
            stem, (None, None, None, None)
        )
        if doi:
            doi_key = _normalise_key(_normalise_doi(doi))
            if doi_key and doi_key in by_doi:
                return by_id[by_doi[doi_key]], "doi"
        if title:
            title_key = _normalise_key(title)
            if title_key in title_lookup and title_lookup[title_key] in by_id:
                return by_id[title_lookup[title_key]], "title"
        if authors:
            author_key = _author_year_key(authors, year)
            if author_key and author_key in by_author_year:
                return by_id[by_author_year[author_key]], "first author + year"
        author_key = _stem_author_year_key(stem)
        if author_key and author_key in by_author_year:
            return by_id[by_author_year[author_key]], "first author + year"
        return None, ""

    return match


async def _link_pdfs(
    repo: MetaRepository,
    session_factory: async_sessionmaker,
    project_path: str,
    papers: list[Path],
    hits: list[Any],
    extraction_identity: dict[str, tuple[Any, Any, Any, Any]],
    resolved: dict[str, int],
    title_lookup: dict[str, int],
    documents: dict[int, Any],
    errors: list[str],
) -> tuple[int, int]:
    """Link each ``Papers/*.pdf`` to one hit; returns ``(linked, unmatched)``.

    A hit is linked at most once per import (a PDF already attached is never
    re-imported, and two file-name variants of the same paper cannot both
    attach to one hit).
    """
    match = _pdf_matcher(hits, resolved, title_lookup, extraction_identity)
    linked = 0
    unmatched = 0
    linked_ids: set[int] = set()
    for pdf in papers:
        hit, _reason = match(pdf.stem)
        if hit is None:
            unmatched += 1
            errors.append(f"{pdf.name}: no matching study")
            continue
        existing = documents.get(hit.hit_id)
        if hit.hit_id in linked_ids or (
            existing is not None and existing.source_id is not None
        ):
            continue
        try:
            await _store_pdf(
                project_path,
                session_factory,
                repo,
                hit,
                content=pdf.read_bytes(),
                filename=pdf.name,
            )
        except (OSError, ValueError) as err:
            errors.append(f"{pdf.name}: {err}")
            continue
        linked_ids.add(hit.hit_id)
        linked += 1
    return linked, unmatched


# ── import ────────────────────────────────────────────────────────────────

async def import_codingstudio(
    session_factory: async_sessionmaker,
    project_path: str,
    folder: str,
    *,
    owner: str = "codingstudio",
    run_id: str = "codingstudio",
    import_hits: bool = True,
    import_screening: bool = True,
    import_extractions: bool = True,
    link_pdfs: bool = True,
    scheme_name: str = "codingstudio",
) -> dict[str, Any]:
    """Import a CodingStudio folder into the open project (idempotent).

    ``session_factory``/``project_path`` come from the active ``ProjectService``.
    Raises ``ValueError`` when the folder is missing or carries neither a
    workbook nor extraction files; every other problem is collected in
    ``errors`` rather than aborting the run.
    """
    scan = _scan_folder(folder)
    root: Path = scan["root"]
    if not root.is_dir():
        raise ValueError(f"CodingStudio folder not found: {folder}")
    workbook: Path = scan["workbook"]
    has_workbook = workbook.is_file()
    if not has_workbook and not scan["extractions"]:
        raise ValueError(
            f"'{folder}' is unusable: it has no {WORKBOOK_NAME} and no Extraction/*.json"
        )

    errors: list[str] = []

    schema: dict[str, Any] | None = None
    if scan["schema"].is_file():
        try:
            loaded = _read_json(scan["schema"])
            schema = loaded if isinstance(loaded, dict) else None
        except (OSError, ValueError) as err:
            errors.append(f"schema.json: {err}")

    async with session_factory() as session:
        repo = MetaRepository(session)

        # 1. criteria (always rewritten — import is the source of truth).
        criteria = _criteria_payload(_intervention_options(schema))
        await repo.replace_criteria(criteria)

        # 2. scheme (guard: never clobber an existing scheme of the same name).
        existing_scheme = await repo.get_scheme(name=scheme_name)
        if existing_scheme is not None:
            try:
                codebook = normalise_codebook(
                    json.loads(existing_scheme.codebook) if existing_scheme.codebook else {}
                )
            except (TypeError, ValueError):
                codebook = normalise_codebook({})
        else:
            raw_book = schema_codebook(schema) if schema is not None else minimal_codebook()
            codebook = normalise_codebook(raw_book)
            await repo.add_scheme(
                name=scheme_name,
                label="CodingStudio",
                description=f"Imported from {folder}",
                codebook=json.dumps(codebook, ensure_ascii=False),
            )
        roles = role_map(all_fields(codebook))

        # 3. hits (workbook rows, or one per extraction file as a fallback).
        if has_workbook:
            try:
                studies = await asyncio.to_thread(_workbook_studies, workbook)
            except (OSError, ValueError) as err:
                raise ValueError(f"{WORKBOOK_NAME} could not be read: {err}") from err
        else:
            studies = _studies_from_extractions(scan["extractions"], codebook)

        sidecar = _load_sidecar(project_path)
        all_hits = await repo.list_hits()
        lookup = _hits_lookup(all_hits)
        title_lookup = _title_lookup(all_hits)
        existing_ids = {hit.hit_id for hit in all_hits}

        hits_imported = 0
        hits_skipped = 0
        if import_hits and studies:
            # A study already mapped to a live hit (through the sidecar, which
            # is keyed by its ``an`` or ``row-<worksheet row>``) is not
            # re-inserted: rows with a title but neither a DOI nor authors have
            # no ``add_hits`` dedupe key, so re-imports would otherwise append
            # a fresh copy of each on every run.
            known = {key for key, hit_id in sidecar.items() if hit_id in existing_ids}
            rows = [
                {**study["hit"], "search_run": run_id}
                for study in studies
                if str(study["key"]) not in known
            ]
            if rows:
                hits_imported = await repo.add_hits(rows)
            hits_skipped = len(rows) - hits_imported
            if hits_imported:
                all_hits = await repo.list_hits()
                lookup = _hits_lookup(all_hits)
                title_lookup = _title_lookup(all_hits)
                existing_ids = {hit.hit_id for hit in all_hits}

        resolved = _resolve_studies(studies, lookup, title_lookup, sidecar, existing_ids)
        for key, hit_id in resolved.items():
            sidecar[key] = hit_id
        for study in studies:
            key = str(study["key"])
            if key not in resolved:
                errors.append(f"study '{key}' could not be matched to a hit")
        _save_sidecar(project_path, sidecar, folder)

        # 4. extraction identity (key -> DOI/title/authors) is needed both to
        #    resolve per-coder full-text rows (keyed by study key, not ``an``)
        #    and to match extraction files to hits below.
        extraction_identity: dict[str, tuple[Any, Any, Any, Any]] = {}
        if import_screening or import_extractions or link_pdfs:
            extraction_identity = await _extraction_identity(scan, codebook)

        # 5. screening: abstract verdicts from Screening/*.json, per-coder and
        #    workbook full text, and the workbook's C1..C4 as the base rater.
        screening_rows = 0
        screening_by_stage: dict[str, int] = {}
        if import_screening:
            screening_rows, screening_by_stage = await _import_screening(
                repo,
                scan,
                studies,
                resolved,
                criteria,
                owner,
                run_id,
                has_workbook,
                errors,
                extraction_identity,
            )

        # 6. extractions: Extraction/<Key>.json (+ Codings/<Key> - <coder>.json).
        extractions_imported = 0
        if import_extractions:
            extractions_imported += await _import_extractions(
                repo,
                scan,
                codebook,
                roles,
                resolved,
                lookup,
                title_lookup,
                owner,
                scheme_name,
                extraction_identity,
                errors,
            )

        # 7. PDFs by key -> hit (study key, DOI, title, first author + year).
        pdfs_linked = 0
        pdfs_unmatched = 0
        if link_pdfs and scan["papers"]:
            documents = {doc.hit_id: doc for doc in await repo.list_documents()}
            pdfs_linked, pdfs_unmatched = await _link_pdfs(
                repo,
                session_factory,
                project_path,
                scan["papers"],
                all_hits,
                extraction_identity,
                resolved,
                title_lookup,
                documents,
                errors,
            )

    return {
        "ok": True,
        "scheme": scheme_name,
        "hits_imported": hits_imported,
        "hits_skipped": hits_skipped,
        "screening_rows": screening_rows,
        "screening_by_stage": screening_by_stage,
        "extractions_imported": extractions_imported,
        "pdfs_linked": pdfs_linked,
        "pdfs_unmatched": pdfs_unmatched,
        "errors": errors[:MAX_ERRORS],
        "errors_truncated": max(0, len(errors) - MAX_ERRORS),
    }


#: Owner of the workbook's own verdicts (the reference treats the workbook as
#: the AI/base rater; coder files are the human raters).
BASE_RATER = "LLM/base"

#: One batched upsert per (hit, owner, run, stage). ``created`` is left as-is
#: on conflict so re-imports only refresh the verdict payload and ``updated``.
_SCREENING_UPSERT = text(
    "INSERT INTO meta_screening (hit_id, owner, run_id, verdicts, intervention_type, "
    "rationale, origin, model, prompt_hash, stage, reason, created, updated) "
    "VALUES (:hit_id, :owner, :run_id, :verdicts, :intervention_type, :rationale, "
    ":origin, :model, :prompt_hash, :stage, :reason, :created, :updated) "
    "ON CONFLICT(hit_id, owner, run_id, stage) DO UPDATE SET "
    "verdicts=excluded.verdicts, intervention_type=excluded.intervention_type, "
    "rationale=excluded.rationale, origin=excluded.origin, model=excluded.model, "
    "prompt_hash=excluded.prompt_hash, reason=excluded.reason, updated=excluded.updated"
)


def _screening_row(
    hit_id: int,
    owner: str,
    run_id: str,
    verdicts: dict[str, Any],
    *,
    stage: str,
    intervention: str | None = None,
    rationale: str | None = None,
    reason: str = "",
) -> dict[str, Any]:
    now = _now()
    return {
        "hit_id": hit_id,
        "owner": owner,
        "run_id": run_id,
        "verdicts": json.dumps(verdicts, ensure_ascii=False),
        "intervention_type": intervention,
        "rationale": rationale,
        "origin": "import",
        "model": None,
        "prompt_hash": None,
        "stage": stage,
        "reason": reason,
        "created": now,
        "updated": now,
    }


async def _write_screening_rows(
    repo: MetaRepository, rows: list[dict[str, Any]]
) -> tuple[int, dict[str, int]]:
    """Batch the screening writes: one ``executemany`` per chunk.

    Building ~30k rows through ``upsert_screening`` would be one SELECT +
    INSERT/UPDATE + commit per row; an ``INSERT … ON CONFLICT DO UPDATE``
    executemany keeps a full-corpus import fast. Rows are deduplicated by the
    unique ``(hit_id, owner, run_id, stage)`` key (last wins) because SQLite
    rejects a batch that touches the same key twice. Returns the number of
    rows stored and the per-stage counts **after** that dedupe, so the caller
    can report counts that match the stored rows rather than the attempts.
    """
    if not rows:
        return 0, {}
    deduped: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        deduped[(row["hit_id"], row["owner"], row["run_id"], row["stage"])] = row
    batch = list(deduped.values())
    by_stage: dict[str, int] = {}
    for row in batch:
        stage = str(row["stage"])
        by_stage[stage] = by_stage.get(stage, 0) + 1
    chunk = 1000
    for start in range(0, len(batch), chunk):
        await repo.session.execute(_SCREENING_UPSERT, batch[start : start + chunk])
        await repo.session.commit()
    return len(batch), by_stage


def _study_resolver(
    studies: list[dict[str, Any]],
    extraction_identity: dict[str, tuple[Any, Any, Any, Any]],
):
    """Resolve a screening entry key to a workbook/extraction study.

    An abstract ``Screening/<coder>.json`` is keyed by the workbook ``an``;
    ``Screening/fulltext - <coder>.json`` is keyed by the **study key** (the
    ``Extraction/<Key>.json`` stem) and carries no ``an``. The resolver tries
    the accession, the study key, the extraction identity (DOI/title) and
    finally a normalised-title match so per-coder full-text rows import.
    """
    by_an: dict[str, dict[str, Any]] = {}
    by_key: dict[str, dict[str, Any]] = {}
    by_identity: dict[str, dict[str, Any]] = {}
    by_title: dict[str, dict[str, Any]] = {}
    by_doi: dict[str, dict[str, Any]] = {}
    for study in studies:
        an = _clean(study.get("an"))
        if an:
            by_an.setdefault(an, study)
        by_key.setdefault(str(study.get("key")), study)
        identity = _identity_key(study.get("doi"), study.get("title"), study.get("authors"))
        if identity:
            by_identity.setdefault(identity, study)
        title = _normalise_key(study.get("title"))
        if title:
            by_title.setdefault(title, study)
        doi = _normalise_key(_normalise_doi(study.get("doi")))
        if doi:
            by_doi.setdefault(doi, study)

    def resolve(accession: Any, entry: Any) -> dict[str, Any] | None:
        key = _clean(accession)
        study = by_an.get(key)
        if study is None and isinstance(entry, dict):
            study = by_an.get(_clean(entry.get("an")))
        if study is None:
            study = by_key.get(key)
        if study is None:
            doi, title, authors, _year = extraction_identity.get(key, (None, None, None, None))
            identity = _identity_key(doi, title, authors)
            if identity:
                study = by_identity.get(identity)
            if study is None and title:
                study = by_title.get(_normalise_key(title))
            if study is None and doi:
                study = by_doi.get(_normalise_key(_normalise_doi(doi)))
        if study is None:
            study = by_title.get(_normalise_key(key))
        return study

    return resolve


async def _import_screening(
    repo: MetaRepository,
    scan: dict[str, Any],
    studies: list[dict[str, Any]],
    resolved: dict[str, int],
    criteria: list[dict[str, Any]],
    owner: str,
    run_id: str,
    has_workbook: bool,
    errors: list[str],
    extraction_identity: dict[str, tuple[Any, Any, Any, Any]],
) -> tuple[int, dict[str, int]]:
    rows: list[dict[str, Any]] = []
    resolve = _study_resolver(studies, extraction_identity)

    for coder, stage, path in _screening_files(scan["screenings"]):
        try:
            data = _read_json(path)
        except (OSError, ValueError) as err:
            errors.append(f"{path.name}: {err}")
            continue
        if not isinstance(data, dict):
            continue
        for accession, entry in data.items():
            if not isinstance(entry, dict):
                continue
            study = resolve(accession, entry)
            if study is None:
                errors.append(f"{path.name}: no study for key '{accession}'")
                continue
            hit_id = resolved.get(str(study["key"]))
            if hit_id is None:
                continue
            if stage == "fulltext":
                verdicts, reason = _fulltext_verdicts(entry, criteria)
                intervention = None
                rationale = _clean(entry.get("note"))
            else:
                verdicts, intervention, rationale = _abstract_verdicts(entry)
                reason = ""
            if not verdicts:
                continue
            rows.append(
                _screening_row(
                    hit_id,
                    coder,
                    run_id,
                    verdicts,
                    stage=stage,
                    intervention=intervention,
                    rationale=rationale or None,
                    reason=reason,
                )
            )

    if has_workbook:
        for study in studies:
            hit_id = resolved.get(str(study["key"]))
            if hit_id is None:
                continue
            # The workbook's C1..C4 + Intervention columns are the base rater's
            # abstract verdicts; its explicit decision is stored as the reason.
            block = study.get("abstract")
            if isinstance(block, dict):
                verdicts, intervention, rationale = _abstract_verdicts(block)
                if verdicts:
                    decision = _clean(study.get("decision"))
                    rows.append(
                        _screening_row(
                            hit_id,
                            BASE_RATER,
                            run_id,
                            verdicts,
                            stage="abstract",
                            intervention=intervention,
                            rationale=rationale or None,
                            reason=decision,
                        )
                    )
            fulltext = study.get("fulltext")
            if fulltext:
                verdicts, reason = _fulltext_verdicts(fulltext, criteria)
                if verdicts:
                    rows.append(
                        _screening_row(
                            hit_id,
                            owner,
                            run_id,
                            verdicts,
                            stage="fulltext",
                            reason=reason,
                        )
                    )

    written, counts = await _write_screening_rows(repo, rows)
    # Report the deduped, post-conflict counts so they are consistent with
    # ``screening_rows`` (``abstract + fulltext == screening_rows``).
    by_stage = {stage: counts.get(stage, 0) for stage in ("abstract", "fulltext")}
    return written, by_stage


async def _extraction_identity(
    scan: dict[str, Any], codebook: dict[str, Any]
) -> dict[str, tuple[Any, Any, Any, Any]]:
    roles = role_map(all_fields(codebook))
    identity: dict[str, tuple[Any, Any, Any, Any]] = {}
    for path in scan["extractions"]:
        try:
            data = _read_json(path)
        except (OSError, ValueError):
            continue
        records, _ = records_from_response(data, codebook)
        identity[path.stem] = _records_identity(records, roles)
    for path in _coder_paths(scan["codings"]):
        key, _, _coder = _coding_parts(path)
        if not key or key in identity:
            continue
        try:
            data = _read_json(path)
        except (OSError, ValueError):
            continue
        records, _ = records_from_response(data, codebook)
        identity[key] = _records_identity(records, roles)
    return identity


def _coding_parts(path: Path) -> tuple[str, str, str]:
    """``(key, coder, raw)`` from a ``Codings/<Key> - <coder>.json`` name."""
    stem = path.stem
    if " - " in stem:
        key, coder = stem.rsplit(" - ", 1)
        return key.strip(), coder.strip(), stem
    return stem.strip(), "", stem


#: ``Codings/`` names that are not a plain extra coder's paper coding: the
#: harmonizer override files and the Space-Bunny variant/route/restored drafts.
_SKIPPED_CODING_MARKERS = (
    " - harmonizer - ",
    " - SB (",
    " (route ",
    " (restored)",
)


def _is_coder_file(path: Path) -> bool:
    stem = path.stem
    return not any(marker in stem for marker in _SKIPPED_CODING_MARKERS)


def _coder_paths(paths: list[Path]) -> list[Path]:
    """Extra-coder ``Codings/*.json`` paths, minus harmonizer/SB drafts."""
    return [path for path in paths if _is_coder_file(path)]


async def _import_extractions(
    repo: MetaRepository,
    scan: dict[str, Any],
    codebook: dict[str, Any],
    roles: dict[str, str],
    resolved: dict[str, int],
    lookup: dict[str, int],
    title_lookup: dict[str, int],
    owner: str,
    scheme_name: str,
    extraction_identity: dict[str, tuple[Any, Any, Any, Any]],
    errors: list[str],
) -> int:
    sources: list[tuple[str, Path, str]] = [
        (path.stem, path, owner) for path in scan["extractions"]
    ]
    for path in _coder_paths(scan["codings"]):
        key, coder, _raw = _coding_parts(path)
        sources.append((key, path, coder or owner))

    written = 0
    for key, path, coder in sources:
        try:
            data = _read_json(path)
        except (OSError, ValueError) as err:
            errors.append(f"{path.name}: {err}")
            continue
        records, missing = records_from_response(data, codebook)
        if not records:
            errors.append(f"{path.name}: no study records found")
            continue
        hit_id = resolved.get(key)
        if hit_id is None:
            doi, title, authors, _year = extraction_identity.get(key, (None, None, None, None))
            hit_id = lookup.get(_identity_key(doi, title, authors) or "")
        if hit_id is None:
            _doi, title, _authors, _y = extraction_identity.get(key, (None, None, None, None))
            hit_id = title_lookup.get(_normalise_key(title))
        if hit_id is None:
            errors.append(f"{path.name}: no matching hit")
            continue
        flat = derive_dv_metrics(records, roles)
        missing_reason = f"{len(missing)} DV comparisons lack full statistics" if missing else None
        await repo.upsert_extraction(
            hit_id=hit_id,
            owner=coder,
            scheme=scheme_name,
            study_vars={"records": records},
            dv_metrics=flat,
            status="extracted",
            missing_reason=missing_reason,
        )
        written += 1
    return written
